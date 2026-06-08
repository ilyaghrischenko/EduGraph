from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def source_citations(chunks: list[str]) -> set[int]:
    numbers: set[int] = set()
    for chunk in chunks:
        for inside in re.findall(r"\[([^\]]+)\]", chunk):
            if not re.fullmatch(r"[\d,; –-]+", inside):
                continue
            for part in re.split(r"[,;]", inside):
                part = part.strip()
                match = re.fullmatch(r"(\d+)\s*[–-]\s*(\d+)", part)
                if match:
                    numbers.update(range(int(match.group(1)), int(match.group(2)) + 1))
                elif part.isdigit() and int(part) <= 20:
                    numbers.add(int(part))
    return numbers


def previous_paragraph_text(body: list, index: int) -> str:
    for element in reversed(body[:index]):
        if element.tag == qn("w:p"):
            text = "".join(node.text or "" for node in element.xpath(".//w:t")).strip()
            if text:
                return text
    return ""


def main() -> None:
    path = Path(sys.argv[1])
    document = Document(path)
    chunks = [paragraph.text for paragraph in document.paragraphs]
    chunks.extend(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    text = "\n".join(chunks)
    citations = source_citations(chunks)
    print(
        f"paragraphs={len(document.paragraphs)} tables={len(document.tables)} "
        f"images={len(document.inline_shapes)}"
    )
    print(f"source_citations={sorted(citations)}")
    print(f"missing_source_citations={sorted(set(range(1, 21)) - citations)}")
    print(f"economic_mentions={len(re.findall('економ', text, re.I))}")
    print(f"practice_mentions={len(re.findall('переддиплом', text, re.I))}")
    print(
        "forbidden_heading_3_10="
        + str(
            sum(
                paragraph.text.strip().startswith("3.10 ")
                for paragraph in document.paragraphs
            )
        )
    )
    print(
        "figure_captions="
        + str(sum(p.text.strip().startswith("Рисунок ") for p in document.paragraphs))
    )
    print(
        "table_captions="
        + str(sum(p.text.strip().startswith("Таблиця ") for p in document.paragraphs))
    )

    body = list(document._body._element)
    uncaptained: list[tuple[int, str]] = []
    table_number = 0
    for index, element in enumerate(body):
        if element.tag != qn("w:tbl"):
            continue
        table_number += 1
        previous = previous_paragraph_text(body, index)
        if table_number > 4 and not previous.startswith("Таблиця "):
            uncaptained.append((table_number, previous))
    print(f"uncaptioned_content_tables={uncaptained}")


if __name__ == "__main__":
    main()
