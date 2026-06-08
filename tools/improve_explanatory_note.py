from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
INPUT_DOCX = ROOT / "outputs" / "Пояснювальна_записка_EduGraph_Грищенко_Ілля_КІ_22_1_розширена.docx"
OUTPUT_DOCX = ROOT / "outputs" / "Пояснювальна_записка_EduGraph_Грищенко_Ілля_КІ_22_1_вдосконалена.docx"
ASSET_DIR = ROOT / "outputs" / "edugraph_note_assets"
SCREENSHOT_DIR = Path("/Users/admin/Library/Containers/me.damir.dropover-mac/Data/tmp/Screenshots")
SCREENSHOT_OUT = ASSET_DIR / "screenshots_clean"
FONT_PATH = Path("/System/Library/Fonts/Supplemental/Times New Roman.ttf")


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(FONT_PATH), size=size)


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font_obj: ImageFont.FreeTypeFont, max_width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    line = ""
    for word in words:
        test = word if not line else f"{line} {word}"
        if draw.textbbox((0, 0), test, font=font_obj)[2] <= max_width:
            line = test
        else:
            if line:
                lines.append(line)
            line = word
    if line:
        lines.append(line)
    return lines


def draw_box(
    draw: ImageDraw.ImageDraw,
    xy: tuple[int, int, int, int],
    text: str,
    fill: str,
    outline: str,
    font_obj: ImageFont.FreeTypeFont,
    text_fill: str = "#1f2937",
) -> None:
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=22, fill=fill, outline=outline, width=3)
    lines = wrap_text(draw, text, font_obj, x2 - x1 - 28)
    total_h = len(lines) * (font_obj.size + 6)
    y = y1 + ((y2 - y1) - total_h) // 2
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font_obj)
        x = x1 + ((x2 - x1) - (bbox[2] - bbox[0])) // 2
        draw.text((x, y), line, fill=text_fill, font=font_obj)
        y += font_obj.size + 6


def draw_arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], fill: str = "#334155") -> None:
    draw.line([start, end], fill=fill, width=3)
    ex, ey = end
    sx, sy = start
    if abs(ex - sx) >= abs(ey - sy):
        direction = 1 if ex > sx else -1
        points = [(ex, ey), (ex - 16 * direction, ey - 8), (ex - 16 * direction, ey + 8)]
    else:
        direction = 1 if ey > sy else -1
        points = [(ex, ey), (ex - 8, ey - 16 * direction), (ex + 8, ey - 16 * direction)]
    draw.polygon(points, fill=fill)


def create_graph_schema(path: Path) -> None:
    img = Image.new("RGB", (1600, 900), "#ffffff")
    draw = ImageDraw.Draw(img)
    title_font = font(42, bold=True)
    node_font = font(28)
    small_font = font(24)
    draw.text((70, 42), "Структура графового подання результатів пошуку EduGraph", fill="#0f172a", font=title_font)
    draw.text((72, 98), "Runtime-граф формується у frontend на основі папок Google Drive та результатів семантичного пошуку.", fill="#475569", font=small_font)

    root = (115, 245, 355, 345)
    trunk = (495, 245, 735, 345)
    folder_a = (880, 150, 1230, 250)
    folder_b = (880, 340, 1230, 440)
    doc_a = (1315, 120, 1535, 230)
    doc_b = (1315, 300, 1535, 410)
    modal = (1130, 610, 1535, 760)

    draw_box(draw, root, "root\nG7", "#d1fae5", "#059669", node_font)
    draw_box(draw, trunk, "trunk\nтехнічний вузол", "#f1f5f9", "#64748b", node_font)
    draw_box(draw, folder_a, "folder\nName, Link, IsMain", "#dcfce7", "#16a34a", node_font)
    draw_box(draw, folder_b, "folder\nName, Link", "#dcfce7", "#16a34a", node_font)
    draw_box(draw, doc_a, "document\ntitle, url,\ncontent, chunkId", "#e0f2fe", "#0284c7", small_font)
    draw_box(draw, doc_b, "document\nfolderName,\nscore", "#e0f2fe", "#0284c7", small_font)
    draw_box(draw, modal, "DocumentModal\nфрагмент документа та посилання на оригінал", "#fef9c3", "#ca8a04", small_font)

    draw_arrow(draw, (355, 295), (495, 295))
    draw_arrow(draw, (735, 295), (880, 200))
    draw_arrow(draw, (735, 295), (880, 390))
    draw_arrow(draw, (1230, 200), (1315, 175), "#0284c7")
    draw_arrow(draw, (1230, 390), (1315, 355), "#0284c7")
    draw_arrow(draw, (1420, 410), (1325, 610), "#ca8a04")

    legend_x = 90
    legend_y = 575
    draw.text((legend_x, legend_y), "Правила побудови", fill="#0f172a", font=font(32, bold=True))
    rules = [
        "Папки завантажуються з endpoint-а /api/google-drive/folders.",
        "Документи додаються після пошуку та прив'язуються до папки за folderName.",
        "Якщо папка не знайдена, документ прив'язується до root-вузла.",
        "Клік по folder відкриває Google Drive, клік по document відкриває modal.",
    ]
    y = legend_y + 48
    for rule in rules:
        draw.ellipse((legend_x, y + 8, legend_x + 12, y + 20), fill="#059669")
        draw.text((legend_x + 26, y), rule, fill="#334155", font=small_font)
        y += 42

    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path)


def create_application_state_schema(path: Path) -> None:
    img = Image.new("RGB", (1600, 900), "#ffffff")
    draw = ImageDraw.Draw(img)
    draw.text((90, 55), "Діаграма станів заявки на реєстрацію", fill="#0f172a", font=font(44, bold=True))
    draw.text((92, 115), "Стан заявки змінюється тільки після рішення уповноваженого користувача.", fill="#475569", font=font(26))
    pending = (150, 330, 500, 470)
    approved = (690, 190, 1080, 330)
    rejected = (690, 500, 1080, 640)
    user = (1180, 190, 1490, 330)
    draw_box(draw, pending, "Pending\nочікує рішення", "#fef9c3", "#ca8a04", font(30))
    draw_box(draw, approved, "Approved\nзаявку схвалено", "#dcfce7", "#16a34a", font(30))
    draw_box(draw, rejected, "Rejected\nзаявку відхилено", "#fee2e2", "#dc2626", font(30))
    draw_box(draw, user, "User\nстворено обліковий запис", "#e0f2fe", "#0284c7", font(30))
    draw_arrow(draw, (500, 380), (690, 260), "#16a34a")
    draw_arrow(draw, (500, 420), (690, 570), "#dc2626")
    draw_arrow(draw, (1080, 260), (1180, 260), "#0284c7")
    draw.text((555, 275), "Схвалити", fill="#166534", font=font(26))
    draw.text((555, 545), "Відхилити", fill="#991b1b", font=font(26))
    draw.text((1105, 215), "Identity", fill="#0369a1", font=font(26))
    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path)


def prepare_screenshots() -> dict[str, Path]:
    SCREENSHOT_OUT.mkdir(parents=True, exist_ok=True)
    mapping = {
        "login": "Снимок экрана — 2026-05-29 в 13.42.32.png",
        "signup_student": "Снимок экрана — 2026-05-29 в 13.42.39.png",
        "search_initial": "Снимок экрана — 2026-05-29 в 13.42.58.png",
        "search_results": "Снимок экрана — 2026-05-29 в 13.43.40.png",
        "document_modal": "Снимок экрана — 2026-05-29 в 13.44.05.png",
        "teacher_search": "Снимок экрана — 2026-05-29 в 13.45.07.png",
        "applications": "Снимок экрана — 2026-05-29 в 13.45.59.png",
        "students": "Снимок экрана — 2026-05-29 в 13.46.07.png",
        "teachers": "Снимок экрана — 2026-05-29 в 13.47.12.png",
        "admins": "Снимок экрана — 2026-05-29 в 13.47.31.png",
        "superadmin_nav": "Снимок экрана — 2026-05-29 в 13.47.27.png",
    }
    out: dict[str, Path] = {}
    for key, name in mapping.items():
        src = SCREENSHOT_DIR / name
        img = Image.open(src).convert("RGB")
        width, height = img.size
        crop_top = 175
        crop_bottom = min(height, 2050)
        cropped = img.crop((0, crop_top, width, crop_bottom))
        cropped.thumbnail((2200, 1300), Image.Resampling.LANCZOS)
        dst = SCREENSHOT_OUT / f"{key}.png"
        cropped.save(dst, optimize=True)
        out[key] = dst
    return out


def set_run_font(run, size_pt: float | None = None, bold: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    if size_pt:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), "Times New Roman")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width_cm: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def format_table(table, widths_cm: Sequence[float] | None = None) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            if widths_cm and col_idx < len(widths_cm):
                set_cell_width(cell, widths_cm[col_idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.first_line_indent = None
                paragraph.paragraph_format.line_spacing = 1.15
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    set_run_font(run, 12)
            if row_idx == 0:
                set_cell_shading(cell, "E2F0D9")
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        set_run_font(run, 12, bold=True)


def paragraph_text(element) -> str:
    return "".join(node.text or "" for node in element.xpath(".//w:t")).strip()


def body(doc: Document):
    return doc._body._element


def body_children(doc: Document) -> list:
    return list(body(doc))


def find_block(doc: Document, text: str, occurrence: int = 1, start: int = 0) -> int:
    count = 0
    for idx, element in enumerate(body_children(doc)[start:], start=start):
        if paragraph_text(element) == text:
            count += 1
            if count == occurrence:
                return idx
    raise ValueError(f"Block not found: {text!r} occurrence={occurrence}")


def find_block_after(doc: Document, text: str, after_idx: int) -> int:
    return find_block(doc, text, occurrence=1, start=after_idx + 1)


def find_project_block(doc: Document, text: str) -> int:
    project_start = find_block(doc, "3 ПРОЄКТНА ЧАСТИНА", occurrence=2)
    return find_block(doc, text, occurrence=1, start=project_start)


def fragment_start(doc: Document) -> int:
    children = body_children(doc)
    if children and children[-1].tag.endswith("sectPr"):
        return len(children) - 1
    return len(children)


def collect_fragment(doc: Document, start_idx: int) -> list:
    b = body(doc)
    children = body_children(doc)
    end_idx = len(children) - 1 if children and children[-1].tag.endswith("sectPr") else len(children)
    fragment = children[start_idx:end_idx]
    for element in fragment:
        b.remove(element)
    return fragment


def insert_fragment(doc: Document, index: int, fragment: Iterable) -> None:
    b = body(doc)
    for offset, element in enumerate(fragment):
        b.insert(index + offset, element)


def remove_range(doc: Document, start_idx: int, end_idx: int) -> None:
    b = body(doc)
    for element in body_children(doc)[start_idx:end_idx]:
        b.remove(element)


def move_range_before(doc: Document, start_text: str, end_text: str, target_text: str) -> None:
    start_idx = find_block(doc, start_text)
    end_idx = find_block_after(doc, end_text, start_idx - 1)
    fragment = body_children(doc)[start_idx:end_idx]
    b = body(doc)
    for element in fragment:
        b.remove(element)
    target_idx = find_project_block(doc, target_text)
    insert_fragment(doc, target_idx, fragment)


def add_paragraph(doc: Document, text: str = "", style: str | None = None):
    p = doc.add_paragraph(text, style=style)
    if text and style is None:
        p.paragraph_format.first_line_indent = Cm(1.25)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for run in p.runs:
        set_run_font(run, 14)
    return p


def add_heading(doc: Document, text: str, level: int):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(6 if level > 1 else 12)
    p.paragraph_format.space_after = Pt(6)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    for run in p.runs:
        set_run_font(run, 14, bold=True)
    return p


def add_caption(doc: Document, text: str):
    p = doc.add_paragraph(text)
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        set_run_font(run, 14)
    return p


def add_image(doc: Document, path: Path, caption: str, width_cm: float = 16.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width_cm))
    add_caption(doc, caption)


def add_table_with_caption(
    doc: Document,
    caption: str,
    headers: Sequence[str],
    rows: Sequence[Sequence[str]],
    widths_cm: Sequence[float] | None = None,
):
    add_caption(doc, caption)
    table = doc.add_table(rows=1, cols=len(headers))
    header_cells = table.rows[0].cells
    for idx, header in enumerate(headers):
        header_cells[idx].text = header
    for row_values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            cells[idx].text = value
    format_table(table, widths_cm)
    add_paragraph(doc, "")
    return table


def build_section_2(doc: Document, screenshots: dict[str, Path], state_schema: Path) -> list:
    start = fragment_start(doc)
    add_heading(doc, "2.1 Проєктування варіантів використання системи", 2)
    add_paragraph(
        doc,
        "Функціональна модель EduGraph описує не окремі екрани, а повні сценарії роботи "
        "користувачів і зовнішніх сервісів. Межами системи є React-клієнт, ASP.NET Core API, "
        "SQLite-сховище, інтеграція з Google Drive та FastAPI-сервіс семантичного пошуку. "
        "Google Drive API і сервіс семантичного пошуку розглядаються як зовнішні учасники "
        "сценаріїв, оскільки вони виконують спеціалізовані дії поза інтерфейсом користувача.",
    )
    add_table_with_caption(
        doc,
        "Таблиця 2.1 – Актори системи EduGraph",
        ["Актор", "Роль у системі", "Основні операції", "Обмеження"],
        [
            ["Неавторизований користувач", "Кандидат на доступ", "Вхід до системи, подання заявки", "Не має доступу до пошуку та документів"],
            ["Студент", "Основний користувач пошуку", "Пошук, перегляд графа, відкриття фрагмента й оригіналу", "Не керує заявками, користувачами та синхронізацією"],
            ["Викладач", "Користувач з розширеним доступом", "Пошук, заявки, створення студентів, запуск синхронізації", "Не керує адміністраторами"],
            ["Адміністратор", "Керування користувачами й заявками", "Пошук, заявки, студенти, викладачі, синхронізація", "Не створює інших адміністраторів"],
            ["Суперадміністратор", "Повний адміністративний доступ", "Усі адміністративні сценарії, зокрема адміністратори", "Використовується обмеженим колом осіб"],
            ["Google Drive API", "Зовнішнє джерело корпусу", "Повернення папок, файлів, посилань і вмісту", "Не відповідає за авторизацію користувачів EduGraph"],
            ["Сервіс семантичного пошуку", "Зовнішній ML/NLP-компонент", "Індексація фрагментів, vector search, reranking", "Не керує користувачами й ролями"],
        ],
        [3.0, 3.6, 5.6, 4.0],
    )
    add_table_with_caption(
        doc,
        "Таблиця 2.2 – Матриця доступу до варіантів використання",
        ["Варіант використання", "Student", "Teacher", "Admin", "SuperAdmin", "Сервіс"],
        [
            ["Увійти до системи", "+", "+", "+", "+", "ASP.NET Identity"],
            ["Подати заявку на реєстрацію", "до створення ролі", "до створення ролі", "-", "-", "Auth API"],
            ["Переглянути структуру папок", "+", "+", "+", "+", "Google Drive API"],
            ["Виконати семантичний пошук", "+", "+", "+", "+", "FastAPI, LanceDB"],
            ["Переглянути фрагмент документа", "+", "+", "+", "+", "React modal"],
            ["Відкрити оригінал у Google Drive", "+", "+", "+", "+", "Google Drive URL"],
            ["Опрацювати заявки", "-", "+", "+", "+", "SignUpApplications"],
            ["Керувати студентами", "-", "+", "+", "+", "Students API"],
            ["Керувати викладачами", "-", "-", "+", "+", "Teachers API"],
            ["Керувати адміністраторами", "-", "-", "-", "+", "Admins API"],
            ["Запустити синхронізацію Google Drive", "-", "+", "+", "+", "Sync queue"],
            ["Індексувати документи", "-", "-", "-", "-", "Фонова служба, FastAPI"],
        ],
        [4.2, 1.8, 1.8, 1.8, 2.0, 4.2],
    )
    add_paragraph(
        doc,
        "Загальна діаграма варіантів використання відображає, що центральним сценарієм є "
        "пошук документів, але він залежить від попередніх сценаріїв авторизації, синхронізації "
        "корпусу та індексації текстових фрагментів.",
    )
    add_image(doc, ASSET_DIR / "usecases.png", "Рисунок 2.1 – Загальна діаграма варіантів використання EduGraph")

    add_heading(doc, "2.2 Специфікації основних сценаріїв", 2)
    add_paragraph(
        doc,
        "Для основних варіантів використання сформовано текстові специфікації. Вони потрібні "
        "для перевірки повноти функціональних вимог: кожен сценарій має актора, передумови, "
        "основний потік, альтернативи, помилки та результат.",
    )
    specs = [
        (
            "UC-01 Увійти до системи",
            "Авторизований користувач однієї з ролей",
            "Користувач має створений обліковий запис; логін і пароль введені у форму входу.",
            "Система перевіряє дані через Identity, оновлює дату останнього входу, формує JWT з role claim і перенаправляє користувача на маршрут його ролі.",
            "Неправильний пароль, порожні поля, заблокований або відсутній користувач, недійсний токен під час наступних запитів.",
            "Користувач отримує доступ тільки до дозволених сторінок.",
        ),
        (
            "UC-02 Подати заявку на реєстрацію",
            "Неавторизований користувач",
            "Користувач відкрив сторінку реєстрації та обрав тип майбутньої ролі.",
            "Користувач вводить логін, ПІБ, пароль, підтвердження пароля та, для студента, групу. Backend валідовує дані й створює заявку зі статусом Pending.",
            "Логін уже зайнятий, існує активна заявка з таким логіном, пароль не підтверджено, ПІБ не відповідає формату.",
            "Заявка зберігається для подальшого рішення уповноваженої ролі.",
        ),
        (
            "UC-03 Опрацювати заявку на реєстрацію",
            "Викладач, адміністратор або суперадміністратор",
            "У системі є заявка зі статусом Pending.",
            "Користувач відкриває список заявок, переглядає дані й натискає «Схвалити» або «Відхилити». При схваленні створюється User і призначається роль; при відхиленні змінюється статус заявки.",
            "Заявку не знайдено, заявка вже оброблена, створення користувача завершилося помилкою Identity.",
            "Заявка переходить у фінальний стан Approved або Rejected.",
        ),
        (
            "UC-04 Виконати семантичний пошук документів",
            "Будь-який авторизований користувач",
            "Користувач авторизований, структура папок завантажена, індекс документів існує або готовий до пошуку.",
            "Користувач вводить запит, frontend викликає endpoint пошуку, backend валідовує Query, FastAPI створює query embedding, виконує пошук у LanceDB, застосовує CrossEncoder reranking і повертає результати.",
            "Порожній запит, відсутність результатів, недоступність FastAPI, помилка vector store або некоректне посилання Google Drive.",
            "На екрані формується граф із папками та знайденими документами.",
        ),
        (
            "UC-05 Переглянути знайдений фрагмент документа",
            "Користувач, який отримав результати пошуку",
            "У графі є вузол документа з title, content, url і folderName.",
            "Користувач натискає на вузол документа. Frontend відкриває модальне вікно, показує релевантний фрагмент і кнопку переходу до оригіналу в Google Drive.",
            "Документ не має безпечного URL; користувач закриває modal без переходу; фрагмент занадто короткий або містить шум після парсингу.",
            "Дані не змінюються; користувач отримує контекст релевантності результату.",
        ),
        (
            "UC-06 Синхронізувати Google Drive",
            "Уповноважений користувач або фонова служба",
            "Налаштований root folder Google Drive і доступ service account.",
            "Система ставить задачу синхронізації в чергу, отримує папки й документи, витягує текст, оновлює SQLite, видаляє відсутні документи з індексу та відправляє змінені документи до FastAPI.",
            "Google Drive недоступний, root folder не знайдено, файл не має текстового шару, FastAPI повертає помилку індексації.",
            "Корпус і векторний індекс наближаються до актуального стану Google Drive.",
        ),
    ]
    for idx, spec in enumerate(specs, start=3):
        add_table_with_caption(
            doc,
            f"Таблиця 2.{idx} – Специфікація {spec[0]}",
            ["Поле", "Зміст"],
            [
                ["Основний актор", spec[1]],
                ["Передумови", spec[2]],
                ["Основний сценарій", spec[3]],
                ["Альтернативи та помилки", spec[4]],
                ["Постумова", spec[5]],
            ],
            [4.0, 12.0],
        )
    add_image(doc, screenshots["login"], "Рисунок 2.2 – Сторінка входу користувача в систему EduGraph")
    add_image(doc, screenshots["signup_student"], "Рисунок 2.3 – Подання заявки на реєстрацію студентом")
    add_image(doc, screenshots["applications"], "Рисунок 2.4 – Перегляд і оброблення заявок на реєстрацію")

    add_heading(doc, "2.3 Діаграми сценаріїв та взаємодії", 2)
    add_paragraph(
        doc,
        "Діаграми сценаріїв показують, як варіанти використання переходять у взаємодію "
        "компонентів. Для EduGraph важливо окремо показати синхронізацію, пошук і життєвий "
        "цикл заявки, тому що ці процеси виконуються різними частинами системи.",
    )
    add_image(doc, ASSET_DIR / "pipeline.png", "Рисунок 2.5 – Конвеєр синхронізації, індексації та пошуку")
    add_image(doc, ASSET_DIR / "sequence_search.png", "Рисунок 2.6 – Діаграма послідовності семантичного пошуку")
    add_image(doc, state_schema, "Рисунок 2.7 – Діаграма станів заявки на реєстрацію")
    add_table_with_caption(
        doc,
        "Таблиця 2.9 – Зв'язок сценаріїв із реалізацією",
        ["Сценарій", "Frontend", "Backend", "Дані / зовнішній сервіс"],
        [
            ["Вхід", "LoginPage, auth utilities", "Auth/LogIn.cs", "User, JWT"],
            ["Заявка", "SignUpPage", "Auth/SignUp.cs", "SignUpApplication"],
            ["Оброблення заявки", "SignUpApplicationsPage", "Approve/RejectSignUpApplication.cs", "Status: Pending, Approved, Rejected"],
            ["Пошук", "StudentSearchPage", "SearchDocuments.cs", "FastAPI /search, LanceDB"],
            ["Перегляд фрагмента", "DocumentModal", "SearchDocuments.Response", "title, content, url, folderName"],
            ["Синхронізація", "GoogleDriveControls", "SyncGoogleDrive.cs, GoogleDriveSyncService.cs", "Google Drive API, UniversityDocument"],
        ],
        [4.2, 4.0, 4.0, 4.0],
    )
    return collect_fragment(doc, start)


def build_graph_model_fragment(doc: Document, graph_schema: Path) -> list:
    start = fragment_start(doc)
    add_heading(doc, "3.3.5 Модель графового подання результатів пошуку", 3)
    add_paragraph(
        doc,
        "Графове подання в EduGraph є runtime-структурою клієнтського інтерфейсу, а не "
        "окремою графовою базою даних. Воно формується у React після завантаження папок "
        "Google Drive та після отримання результатів семантичного пошуку. Такий підхід "
        "достатній для задачі навігації: користувач бачить, до яких розділів корпусу належать "
        "знайдені документи, але система не вводить окрему онтологічну модель предметної області.",
    )
    add_table_with_caption(
        doc,
        "Таблиця 3.1 – Структура вузлів графового подання",
        ["Тип вузла", "Джерело даних", "Основні поля", "Призначення"],
        [
            ["root", "Константа frontend", "id, title", "Початкова точка графа, що відповідає корпусу G7"],
            ["trunk", "Алгоритм побудови layout", "id, parentId", "Технічний вузол для акуратного розгалуження структури"],
            ["folder", "Endpoint /google-drive/folders", "id, title, url, parentId", "Відображення папки Google Drive і перехід до неї"],
            ["document", "Endpoint /google-drive/search-documents", "title, url, content, folderName, chunkId", "Результат пошуку з релевантним фрагментом"],
        ],
        [3.0, 4.1, 4.4, 4.4],
    )
    add_paragraph(
        doc,
        "Зв'язок документа з папкою в інтерфейсі виконується за значенням folderName. Якщо "
        "відповідна папка знайдена в уже завантаженій структурі, вузол документа приєднується "
        "до неї; якщо ні, документ приєднується до root-вузла. Це запобігає втраті результату "
        "навіть тоді, коли назва папки в індексі не збігається зі списком актуальних папок.",
    )
    add_image(doc, graph_schema, "Рисунок 3.2 – Структура графового подання результатів пошуку EduGraph")
    return collect_fragment(doc, start)


def build_interface_screens_fragment(doc: Document, screenshots: dict[str, Path]) -> list:
    start = fragment_start(doc)
    add_paragraph(
        doc,
        "На рисунках 3.3–3.6 наведено ключові стани інтерфейсу. Вони показують не лише "
        "зовнішній вигляд сторінок, а й логіку взаємодії: завантаження структури папок, "
        "появу документів після пошуку, перегляд релевантного фрагмента та відмінність "
        "робочого сценарію викладача від студентського.",
    )
    add_image(doc, screenshots["search_initial"], "Рисунок 3.3 – Початковий стан графової навігації корпусом документів")
    add_image(doc, screenshots["search_results"], "Рисунок 3.4 – Результати семантичного пошуку у вигляді графу")
    add_image(doc, screenshots["document_modal"], "Рисунок 3.5 – Перегляд релевантного фрагмента документа")
    add_image(doc, screenshots["teacher_search"], "Рисунок 3.6 – Інтерфейс пошуку для користувача з роллю викладача")
    return collect_fragment(doc, start)


def build_appendix_screens_fragment(doc: Document, screenshots: dict[str, Path]) -> list:
    start = fragment_start(doc)
    add_heading(doc, "А.6 Сценарні екрани інтерфейсу EduGraph", 3)
    add_paragraph(
        doc,
        "Сценарні екрани доповнюють основні схеми та показують реалізацію рольової моделі "
        "у користувацькому інтерфейсі. Основний текст містить лише ключові стани, а в додатку "
        "наведено адміністративні сторінки, які підтверджують наявність окремих сценаріїв "
        "для викладача, адміністратора та суперадміністратора.",
    )
    add_image(doc, screenshots["students"], "Рисунок А.6 – Сторінка керування студентами")
    add_image(doc, screenshots["teachers"], "Рисунок А.7 – Сторінка керування викладачами")
    add_image(doc, screenshots["admins"], "Рисунок А.8 – Сторінка керування адміністраторами")
    add_image(doc, screenshots["superadmin_nav"], "Рисунок А.9 – Навігація суперадміністратора в інтерфейсі пошуку")
    return collect_fragment(doc, start)


def build_sources_fragment(doc: Document) -> list:
    start = fragment_start(doc)
    sources = [
        "1. Вихідний код проєкту EduGraph: frontend/edugraph.client, backend/api, backend/vector-search.",
        "2. Методичні вказівки до виконання кваліфікаційної роботи бакалавра зі спеціальності 123 «Комп'ютерна інженерія». Версія 2025.",
        "3. ДСТУ 8302:2015. Інформація та документація. Бібліографічне посилання. Загальні положення та правила складання.",
        "4. Microsoft Learn. Minimal APIs overview. URL: https://learn.microsoft.com/en-us/aspnet/core/fundamentals/minimal-apis",
        "5. Microsoft Learn. Entity Framework Core documentation. URL: https://learn.microsoft.com/en-us/ef/core/",
        "6. Microsoft Learn. ASP.NET Core Identity. URL: https://learn.microsoft.com/en-us/aspnet/core/security/authentication/identity",
        "7. React Documentation. URL: https://react.dev/",
        "8. Vite Documentation. URL: https://vite.dev/",
        "9. FastAPI Documentation. URL: https://fastapi.tiangolo.com/",
        "10. Google for Developers. Google Drive API Documentation. URL: https://developers.google.com/drive/api",
        "11. LanceDB Documentation. Vector Search. URL: https://docs.lancedb.com/search/vector-search",
        "12. Sentence Transformers Documentation. Semantic Search. URL: https://www.sbert.net/examples/sentence_transformer/applications/semantic-search/README.html",
        "13. Manning C. D., Raghavan P., Schütze H. Introduction to Information Retrieval. Cambridge University Press, 2008. URL: https://nlp.stanford.edu/IR-book/",
        "14. Reimers N., Gurevych I. Sentence-BERT: Sentence Embeddings using Siamese BERT-Networks. EMNLP-IJCNLP, 2019. URL: https://arxiv.org/abs/1908.10084",
        "15. Wang L. et al. Text Embeddings by Weakly-Supervised Contrastive Pre-training. 2022. URL: https://arxiv.org/abs/2212.03533",
        "16. Nogueira R., Cho K. Passage Re-ranking with BERT. 2019. URL: https://arxiv.org/abs/1901.04085",
        "17. OpenAPI Specification. URL: https://spec.openapis.org/oas/latest.html",
        "18. Jones M., Bradley J., Sakimura N. JSON Web Token (JWT). RFC 7519. URL: https://www.rfc-editor.org/rfc/rfc7519",
    ]
    for source in sources:
        p = add_paragraph(doc, source)
        p.paragraph_format.first_line_indent = None
        p.paragraph_format.left_indent = Cm(0)
    return collect_fragment(doc, start)


def replace_all_text(doc: Document, replacements: dict[str, str]) -> None:
    for paragraph in doc.paragraphs:
        text = paragraph.text
        new_text = text
        for old, new in replacements.items():
            new_text = new_text.replace(old, new)
        if new_text != text:
            paragraph.text = new_text
            for run in paragraph.runs:
                set_run_font(run, 14, bold=paragraph.style.name.startswith("Heading"))
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    text = paragraph.text
                    new_text = text
                    for old, new in replacements.items():
                        new_text = new_text.replace(old, new)
                    if new_text != text:
                        paragraph.text = new_text
                        for run in paragraph.runs:
                            set_run_font(run, 11)


def clean_draft_phrases(doc: Document) -> None:
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text.startswith("Додаток А містить графічні матеріали"):
            paragraph.text = (
                "Додаток А містить графічні матеріали, які узагальнюють основні рішення "
                "пояснювальної записки: архітектуру системи, конвеєр синхронізації та пошуку, "
                "варіанти використання, послідовність виконання семантичного пошуку, логічну "
                "модель даних і сценарні екрани інтерфейсу EduGraph."
            )
        elif "Після погодження з керівником роботи" in text:
            paragraph.text = (
                "Окремий розділ з охорони праці та безпеки в надзвичайних ситуаціях у цей "
                "документ не включено, оскільки він оформлюється як самостійний документ."
            )
        elif text.startswith("Пояснювальна записка побудована як розширення звіту"):
            paragraph.text = (
                "Пояснювальна записка побудована як розвиток матеріалів звіту з переддипломної "
                "практики. У дипломній роботі ці матеріали деталізовано з погляду вимог, "
                "варіантів використання, структури даних, графового подання результатів, "
                "інтерфейсу та функціональної перевірки системи."
            )
        elif "Оскільки точні ставки, нормативи й коефіцієнти" in text:
            paragraph.text = (
                "Економічно-організаційна частина в цій роботі має розрахунково-обґрунтувальний "
                "характер і показує, як можна оцінити трудомісткість розроблення EduGraph. "
                "Оскільки точні ставки, нормативи й коефіцієнти залежать від методичних вимог "
                "кафедри, у записці подано структуру розрахунку та орієнтовний розподіл робіт."
            )
        elif text.startswith("У фінальній версії цей додаток"):
            paragraph.text = (
                "Цей додаток доповнює основну інструкцію користувача й показує практичний "
                "порядок використання системи після розгортання. Він потрібний для того, щоб "
                "відокремити короткий опис у розділі 3.9 від розгорнутих експлуатаційних "
                "кроків і чек-листів супроводу."
            )
        for run in paragraph.runs:
            set_run_font(run, 14, bold=paragraph.style.name.startswith("Heading"))
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if "Стан у чернетці" in paragraph.text:
                        paragraph.text = paragraph.text.replace("Стан у чернетці", "Стан у записці")
                    if "чернетці" in paragraph.text:
                        paragraph.text = paragraph.text.replace("чернетці", "записці")
                    for run in paragraph.runs:
                        set_run_font(run, 11)


def insert_hard_page_breaks(doc: Document) -> None:
    headings = [
        paragraph
        for paragraph in doc.paragraphs
        if paragraph.text.strip() and paragraph.style.name == "Heading 1"
    ]
    for heading in headings:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = None
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run()
        run.add_break(WD_BREAK.PAGE)
        set_run_font(run, 14)
        element = p._element
        b = body(doc)
        b.remove(element)
        heading._element.addprevious(element)


def apply_document_format(doc: Document) -> None:
    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(3)
        section.right_margin = Cm(1.5)
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)

    for style_name in ["Normal", "Heading 1", "Heading 2", "Heading 3", "List Bullet"]:
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(14)
        if style_name.startswith("Heading"):
            style.font.bold = True
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        style_name = paragraph.style.name
        fmt = paragraph.paragraph_format
        fmt.line_spacing = 1.5
        fmt.space_after = Pt(0)
        if style_name == "Heading 1":
            fmt.first_line_indent = None
            fmt.space_before = Pt(12)
            fmt.space_after = Pt(6)
            fmt.page_break_before = True
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif style_name in {"Heading 2", "Heading 3"}:
            fmt.first_line_indent = None
            fmt.space_before = Pt(6)
            fmt.space_after = Pt(6)
            fmt.page_break_before = False
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif text.startswith("Рисунок") or text.startswith("Таблиця"):
            fmt.first_line_indent = None
            fmt.space_before = Pt(3)
            fmt.space_after = Pt(6)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif not text:
            fmt.first_line_indent = None
        else:
            fmt.first_line_indent = Cm(1.25)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for run in paragraph.runs:
            set_run_font(run, 14, bold=True if style_name.startswith("Heading") else None)

    for table in doc.tables:
        format_table(table)


def main() -> None:
    screenshots = prepare_screenshots()
    graph_schema = ASSET_DIR / "graph_schema.png"
    state_schema = ASSET_DIR / "application_state.png"
    create_graph_schema(graph_schema)
    create_application_state_schema(state_schema)

    doc = Document(INPUT_DOCX)
    replace_all_text(
        doc,
        {
            "2.1 Проєктування UseCases": "2.1 Проєктування варіантів використання системи",
            "2.2 Діаграми сценаріїв": "2.2 Специфікації основних сценаріїв",
            "2.3 Діаграми взаємодії": "2.3 Діаграми сценаріїв та взаємодії",
            "ПЕРЕЛІК ПОСИЛАНЬ": "ПЕРЕЛІК ВИКОРИСТАНИХ ДЖЕРЕЛ",
            "use case-и": "варіанти використання",
            "Use case": "Варіант використання",
            "Use Case": "Варіант використання",
        },
    )
    clean_draft_phrases(doc)

    section2_start = find_block(doc, "2.1 Проєктування варіантів використання системи", occurrence=2)
    section3_start = find_block(doc, "3 ПРОЄКТНА ЧАСТИНА", occurrence=2)
    remove_range(doc, section2_start, section3_start)
    insert_fragment(doc, section2_start, build_section_2(doc, screenshots, state_schema))

    insert_idx = find_project_block(doc, "3.4 Проєктування серверної частини застосунку")
    insert_fragment(doc, insert_idx, build_graph_model_fragment(doc, graph_schema))

    move_range_before(
        doc,
        "3.4.1 Маршрути API та політики доступу",
        "3.5.1 Навігація, маршрути та рольовий frontend",
        "3.5 Проєктування клієнтської частини застосунку",
    )
    move_range_before(
        doc,
        "3.5.1 Навігація, маршрути та рольовий frontend",
        "3.6.1 Алгоритм семантичного пошуку у FastAPI-сервісі",
        "3.6 Опис основних класів, методів і змінних",
    )
    move_range_before(
        doc,
        "3.6.1 Алгоритм семантичного пошуку у FastAPI-сервісі",
        "3.8.1 Деталізована матриця функціональної перевірки",
        "3.7 Структура інтерфейсу користувача",
    )
    move_range_before(
        doc,
        "3.8.1 Деталізована матриця функціональної перевірки",
        "3.10.1 Обмеження поточної реалізації та напрями розвитку",
        "3.9 Інструкція користувача",
    )
    move_range_before(
        doc,
        "3.10.1 Обмеження поточної реалізації та напрями розвитку",
        "5 ЕКОНОМІЧНО-ОРГАНІЗАЦІЙНА ЧАСТИНА",
        "5 ЕКОНОМІЧНО-ОРГАНІЗАЦІЙНА ЧАСТИНА",
    )

    interface_insert_idx = find_project_block(doc, "3.8 Експериментальне тестування спроєктованого застосунку")
    insert_fragment(doc, interface_insert_idx, build_interface_screens_fragment(doc, screenshots))

    sources_start = find_block(doc, "ПЕРЕЛІК ВИКОРИСТАНИХ ДЖЕРЕЛ", occurrence=2)
    appendices_start = find_block(doc, "ДОДАТКИ", occurrence=2)
    remove_range(doc, sources_start + 1, appendices_start)
    insert_fragment(doc, sources_start + 1, build_sources_fragment(doc))

    appendix_insert_idx = find_block(doc, "Додаток Б. Текстова специфікація API та моделей даних")
    insert_fragment(doc, appendix_insert_idx, build_appendix_screens_fragment(doc, screenshots))

    apply_document_format(doc)
    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    main()
